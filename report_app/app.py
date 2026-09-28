"""
report_app.py — Secure, cloud-ready report & export generator for myCompetency Scorecard.

Upload-only: no local file fallbacks, no folder dependencies, no disk writes.
Secrets (Azure OpenAI) read from st.secrets (Streamlit Cloud) or environment variables.
"""

import io
import os
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill

# Load .env from this folder or any parent directory (local dev only)
try:
    from dotenv import load_dotenv, find_dotenv
    load_dotenv(find_dotenv(usecwd=False, raise_error_if_not_found=False))
except ImportError:
    pass

# ============================================================
# CONSTANTS
# ============================================================

DETAILS_SHEET_NAME = "Details"

ALLOWED_BUSINESS_GROUPS = ["Tech_Song", "Tech_Adobe Platform"]

COL_PERSONNEL_NO        = "Personnel No"
COL_EID                 = "Enterpriseid"
COL_LEVEL               = "Management Level"
COL_SKILL_NAME          = "SkillName"
COL_SKILL_TYPE          = "Skill type"
COL_BUSINESS_GROUP      = "Business Group"
COL_PROJECT             = "Project Name"

RESULT_COL_EID          = "Enterpriseid"
RESULT_COL_SKILL        = "SkillName"
RESULT_COL_PROFICIENCY      = "proficiency"
RESULT_COL_PROFICIENCY_DESC = "Proficiency Description"

PROFICIENCY_RANK_MAP  = {"P0": 0, "P1": 1, "P2": 2, "P3": 3, "Expert eligible": 4}
PROFICIENCY_LABEL_MAP = {v: k for k, v in PROFICIENCY_RANK_MAP.items()}


# ============================================================
# SECRETS — st.secrets (Streamlit Cloud) → os.environ fallback
# ============================================================

def _get_secret(key: str, default: str = "") -> str:
    try:
        return st.secrets[key]
    except Exception:
        return os.environ.get(key, default)


# ============================================================
# HELPERS
# ============================================================

def _norm(v):
    return "" if pd.isna(v) else str(v).strip()


def _norm_key(v):
    if pd.isna(v):
        return ""
    t = str(v).strip()
    try:
        return str(int(float(t)))
    except Exception:
        return t


def _norm_prof_desc(v):
    text = _norm(v)
    if text.upper() in ("", "NULL", "NAN", "NONE"):
        return ""
    upper = text.upper()
    rank_upper = {k.upper(): k for k in PROFICIENCY_RANK_MAP}
    if upper in rank_upper:
        return rank_upper[upper]
    if text.isdigit():
        return PROFICIENCY_LABEL_MAP.get(int(text), text)
    try:
        return PROFICIENCY_LABEL_MAP.get(int(float(text)), text)
    except Exception:
        return text


def _target_prof(management_level):
    level = pd.to_numeric(management_level, errors="coerce")
    if pd.isna(level):
        return None
    level = int(level)
    if level in [11, 12]:
        return 2   # P2
    if level <= 10:
        return 3   # P3
    return None


def _safe_sheet(name):
    return re.sub(r'[\\/*?:\[\]]', "_", str(name))[:31]


def _style_workbook(writer, freeze="A2", max_col_width=45):
    for ws in writer.sheets.values():
        ws.freeze_panes = freeze
        for col_cells in ws.columns:
            max_len = max(
                (len(str(c.value)) if c.value else 0 for c in col_cells),
                default=10,
            )
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, max_col_width)


# ============================================================
# DATA BUILDER
# ============================================================

@st.cache_data(show_spinner=False)
def build_data(master_bytes: bytes, result_bytes: bytes, selected_bg: str):
    selected_skill_type = "Primary"

    # ── Master ──────────────────────────────────────────────
    master_df = pd.read_excel(io.BytesIO(master_bytes), sheet_name=DETAILS_SHEET_NAME, header=1)
    master_df.columns = [str(c).strip() for c in master_df.columns]

    master_df[COL_PERSONNEL_NO]   = master_df[COL_PERSONNEL_NO].apply(_norm_key)
    master_df[COL_EID]            = master_df[COL_EID].apply(_norm).str.lower()
    master_df[COL_SKILL_NAME]     = master_df[COL_SKILL_NAME].apply(_norm)
    master_df[COL_SKILL_TYPE]     = master_df[COL_SKILL_TYPE].apply(_norm)
    master_df[COL_BUSINESS_GROUP] = master_df[COL_BUSINESS_GROUP].apply(_norm)
    master_df[COL_PROJECT]        = master_df[COL_PROJECT].fillna("Unmapped").astype(str).str.strip()

    if selected_bg == "All":
        master_df = master_df[
            master_df[COL_BUSINESS_GROUP].str.upper().isin(
                [bg.upper() for bg in ALLOWED_BUSINESS_GROUPS]
            )
        ].copy()
    else:
        master_df = master_df[
            master_df[COL_BUSINESS_GROUP].str.upper() == selected_bg.upper()
        ].copy()

    master_df = master_df[
        master_df[COL_SKILL_TYPE].str.upper() == selected_skill_type.upper()
    ].copy()

    # ── Result (dump) ────────────────────────────────────────
    result_df = pd.read_excel(io.BytesIO(result_bytes))
    result_df.columns = [str(c).strip() for c in result_df.columns]
    result_df[RESULT_COL_EID]   = result_df[RESULT_COL_EID].apply(_norm).str.lower()
    result_df[RESULT_COL_SKILL] = result_df[RESULT_COL_SKILL].apply(_norm)

    if COL_BUSINESS_GROUP in result_df.columns:
        result_df[COL_BUSINESS_GROUP] = result_df[COL_BUSINESS_GROUP].apply(_norm)
        result_df = result_df[
            result_df[COL_BUSINESS_GROUP].str.upper().isin(
                [bg.upper() for bg in ALLOWED_BUSINESS_GROUPS]
            )
        ].copy()

    if COL_SKILL_TYPE in result_df.columns:
        result_df[COL_SKILL_TYPE] = result_df[COL_SKILL_TYPE].apply(_norm)
        result_df = result_df[
            result_df[COL_SKILL_TYPE].str.upper() == selected_skill_type.upper()
        ].copy()

    keep = [RESULT_COL_EID, RESULT_COL_PROFICIENCY]
    if RESULT_COL_PROFICIENCY_DESC in result_df.columns:
        keep.append(RESULT_COL_PROFICIENCY_DESC)
    result_df = result_df[keep].copy()

    result_df = result_df.rename(columns={
        RESULT_COL_EID:              COL_EID,
        RESULT_COL_PROFICIENCY:      "Result Proficiency",
        RESULT_COL_PROFICIENCY_DESC: "Result Proficiency Description",
    })

    result_df = result_df.sort_values("Result Proficiency", ascending=False, na_position="last")
    result_df = result_df.drop_duplicates(subset=[COL_EID], keep="first")

    if "Result Proficiency Description" not in result_df.columns:
        result_df["Result Proficiency Description"] = result_df["Result Proficiency"]
    result_df["Result Proficiency Description"] = (
        result_df["Result Proficiency Description"].apply(_norm_prof_desc)
    )

    # ── Map proficiency onto master (EID lookup, row count never changes) ──
    eid_prof_num  = result_df.set_index(COL_EID)["Result Proficiency"]
    eid_prof_desc = result_df.set_index(COL_EID)["Result Proficiency Description"]
    master_df["Result Proficiency"]             = master_df[COL_EID].map(eid_prof_num)
    master_df["Result Proficiency Description"] = master_df[COL_EID].map(eid_prof_desc)

    merged_df = master_df.copy()

    # ── Competency logic ─────────────────────────────────────
    merged_df["proficiency_desc_clean"] = (
        merged_df["Result Proficiency Description"].astype(str).str.strip()
    )
    merged_df.loc[
        merged_df["proficiency_desc_clean"].str.upper().isin(["NULL", "NAN", "NONE", ""]),
        "proficiency_desc_clean",
    ] = None

    merged_df["proficiency_num"]        = merged_df["proficiency_desc_clean"].map(PROFICIENCY_RANK_MAP)
    merged_df["has_assessment"]         = merged_df["proficiency_num"].notna()
    merged_df["target_proficiency_num"] = merged_df[COL_LEVEL].apply(_target_prof)

    merged_df["meets_target"] = (
        merged_df["has_assessment"]
        & merged_df["target_proficiency_num"].notna()
        & (merged_df["proficiency_num"] >= merged_df["target_proficiency_num"])
    )
    merged_df["below_target"] = (
        merged_df["has_assessment"]
        & merged_df["target_proficiency_num"].notna()
        & (merged_df["proficiency_num"] < merged_df["target_proficiency_num"])
    )

    merged_df["Action Reason"] = "Meeting Target"
    merged_df.loc[~merged_df["has_assessment"], "Action Reason"] = "No Assessment"
    merged_df.loc[merged_df["below_target"],    "Action Reason"] = "Below Target"

    # ── Resource-level aggregation ───────────────────────────
    resource_df = (
        merged_df
        .groupby(COL_PERSONNEL_NO, as_index=False)
        .agg(
            EID                  =(COL_EID,               "first"),
            BusinessGroup        =(COL_BUSINESS_GROUP,    "first"),
            ManagementLevel      =(COL_LEVEL,             "first"),
            Project              =(COL_PROJECT,           "first"),
            SkillName            =(COL_SKILL_NAME,        "first"),
            SkillType            =(COL_SKILL_TYPE,        "first"),
            HasAssessment        =("has_assessment",      "max"),
            MeetingTarget        =("meets_target",        "max"),
            BelowTarget          =("below_target",        "max"),
            ActionReason         =("Action Reason",       "first"),
            ActualProficiencyRank=("proficiency_num",        "max"),
            TargetProficiencyRank=("target_proficiency_num", "first"),
        )
    )
    resource_df["ActualProficiency"] = resource_df["ActualProficiencyRank"].map(PROFICIENCY_LABEL_MAP)
    resource_df["TargetProficiency"] = resource_df["TargetProficiencyRank"].map(PROFICIENCY_LABEL_MAP)

    # ── No Assessment list with Details column ───────────────
    COL_DETAILS = "Details"
    no_asmt_mask = ~merged_df["has_assessment"]

    if COL_DETAILS in merged_df.columns:
        merged_df[COL_DETAILS] = merged_df[COL_DETAILS].fillna("(NA)").astype(str).str.strip()
        no_asmt_df = merged_df[no_asmt_mask][[COL_EID, COL_SKILL_NAME, COL_DETAILS]].copy()
    else:
        no_asmt_df = merged_df[no_asmt_mask][[COL_EID, COL_SKILL_NAME]].copy()
        no_asmt_df[COL_DETAILS] = "(column not found)"

    no_asmt_df = (
        no_asmt_df
        .rename(columns={COL_EID: "EID", COL_SKILL_NAME: "Primary Skill", COL_DETAILS: "Details"})
        .drop_duplicates(subset=["EID"])
        .sort_values("Details")
        .reset_index(drop=True)
    )

    return merged_df, resource_df, no_asmt_df


# ============================================================
# EXPORT BUILDERS  (all in-memory — no disk writes)
# ============================================================

def build_skill_chase_workbook(skill_summary_df, merged_df, resource_df):
    output = io.BytesIO()
    action_order = {"No Assessment": 1, "Below Target": 2, "Meeting Target": 3}

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        skill_summary_df.to_excel(writer, sheet_name="Per Skill Summary", index=False)

        for _, row in skill_summary_df.iterrows():
            skill_name = row[COL_SKILL_NAME]
            df = merged_df[merged_df[COL_SKILL_NAME] == skill_name].copy()
            df = df.merge(
                resource_df[[COL_PERSONNEL_NO, "EID", "Project", "ManagementLevel",
                             "ActionReason", "ActualProficiency", "TargetProficiency"]],
                on=COL_PERSONNEL_NO, how="left",
            ).drop_duplicates(subset=[COL_PERSONNEL_NO])
            df = df[df["ActionReason"].isin(["No Assessment", "Below Target"])].copy()
            if df.empty:
                continue

            df["_ao"] = df["ActionReason"].map(action_order).fillna(999)
            df["_ls"] = pd.to_numeric(df["ManagementLevel"], errors="coerce")
            df = df.sort_values(["_ao", "_ls", COL_PERSONNEL_NO], na_position="last")

            export = df[[COL_PERSONNEL_NO, "EID", "Project", "ManagementLevel",
                         "TargetProficiency", "ActualProficiency", "ActionReason"]].copy()
            export = export.rename(columns={
                COL_PERSONNEL_NO: "Personnel No",
                "ManagementLevel": "Level",
                "TargetProficiency": "Target Proficiency",
                "ActualProficiency": "Actual Proficiency",
                "ActionReason": "Action Reason",
            })
            export.to_excel(writer, sheet_name=_safe_sheet(skill_name), index=False)

        _style_workbook(writer)

    output.seek(0)
    return output


def build_people_lead_excel(summary_df, project_df, resource_df, skill_gap_df, career_df):
    output = io.BytesIO()

    def _grouped(df):
        rows, cols = [], list(df.columns)
        for proj, grp in df.groupby("Project", sort=True):
            hdr = {c: "" for c in cols}
            hdr["Project"] = f"PROJECT: {proj}"
            rows.append(hdr)
            rows.extend(grp.to_dict("records"))
        return pd.DataFrame(rows, columns=cols)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Executive Summary", index=False)
        project_df.to_excel(writer, sheet_name="Project Action Summary", index=False)
        _grouped(resource_df).to_excel(writer, sheet_name="Resource Chase Detail", index=False)
        skill_gap_df.to_excel(writer, sheet_name="Skill Gap Summary", index=False)
        career_df.to_excel(writer, sheet_name="Career Level Health", index=False)

        hdr_fill  = PatternFill(fill_type="solid", fgColor="D9EAF7")
        val_fill  = PatternFill(fill_type="solid", fgColor="F3F8FC")
        proj_fill = PatternFill(fill_type="solid", fgColor="BDD7EE")

        ws_exec = writer.sheets["Executive Summary"]
        for cell in ws_exec[1]:
            cell.font = Font(bold=True, size=12)
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for cell in ws_exec[2]:
            cell.font = Font(bold=True, size=12)
            cell.fill = val_fill
            cell.alignment = Alignment(horizontal="center")

        ws_chase = writer.sheets["Resource Chase Detail"]
        for row in ws_chase.iter_rows(min_row=2):
            val = row[1].value if len(row) > 1 else None
            if isinstance(val, str) and val.startswith("PROJECT:"):
                for cell in row:
                    cell.font = Font(bold=True, size=11)
                    cell.fill = proj_fill

        _style_workbook(writer)

    output.seek(0)
    return output


def build_no_assessment_excel(df):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="No Assessment List", index=False)
        ws = writer.sheets["No Assessment List"]
        ws.freeze_panes = "A2"
        hdr_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for col_cells in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0 for c in col_cells), default=10)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 50)
    output.seek(0)
    return output


# ============================================================
# STREAMLIT APP
# ============================================================

st.set_page_config(
    page_title="myCompetency Report Generator",
    page_icon="📄",
    layout="wide",
)

st.markdown("""
<style>
html, body, [class*="css"] { font-family: "Inter", "Segoe UI", sans-serif; }
#MainMenu, footer { visibility: hidden; }
h1 { font-size: 1.65rem !important; font-weight: 700 !important; }
.metric-row { display:flex; gap:12px; flex-wrap:wrap; margin:1rem 0 1.5rem 0; }
.metric-card {
    flex:1 1 130px; background:#f8fafc; border:1px solid #e2e8f0;
    border-radius:10px; padding:14px 18px; min-width:110px;
}
.metric-card.blue  { border-left:4px solid #3b82f6; }
.metric-card.green { border-left:4px solid #22c55e; }
.metric-card.amber { border-left:4px solid #f59e0b; }
.metric-card.red   { border-left:4px solid #ef4444; }
.metric-card .lbl  { font-size:.7rem; font-weight:600; text-transform:uppercase;
                     letter-spacing:.07em; color:#64748b; margin-bottom:4px; }
.metric-card .val  { font-size:1.6rem; font-weight:700; color:#0f172a; line-height:1; }
.metric-card .sub  { font-size:.7rem; color:#94a3b8; margin-top:3px; }
.section-hdr { margin-top:1.8rem; margin-bottom:.4rem; padding-bottom:.35rem;
               border-bottom:2px solid #e2e8f0; font-size:1rem;
               font-weight:600; color:#1e293b; }
</style>
""", unsafe_allow_html=True)


def _card(label, value, sub="", accent="blue"):
    return (
        f'<div class="metric-card {accent}">'
        f'<div class="lbl">{label}</div>'
        f'<div class="val">{value}</div>'
        + (f'<div class="sub">{sub}</div>' if sub else "")
        + "</div>"
    )


def _section(title, icon=""):
    st.markdown(f'<div class="section-hdr">{icon} {title}</div>', unsafe_allow_html=True)


# ── Sidebar ───────────────────────────────────────────────────────────────────

st.sidebar.title("📄 Report Generator")
st.sidebar.markdown("---")
st.sidebar.markdown("**📂 Upload Source Files**")

master_upload = st.sidebar.file_uploader(
    "1. Master File (.xlsx)", type=["xlsx"], key="master",
    help="MyC Report — Details sheet, header row 2",
)
result_upload = st.sidebar.file_uploader(
    "2. Result / Dump File (.xlsx)", type=["xlsx"], key="result",
    help="Assessment dump with proficiency scores",
)

st.sidebar.markdown("---")
st.sidebar.markdown("**🔍 Scope**")
selected_bg = st.sidebar.selectbox(
    "Business Group",
    ["All"] + ALLOWED_BUSINESS_GROUPS,
    index=0,
)

st.sidebar.markdown("---")
st.sidebar.markdown("**⚙️ Options**")
show_eid = st.sidebar.checkbox("Use EID as identifier in exports", value=True)

# ── Header ────────────────────────────────────────────────────────────────────

st.title("📄 myCompetency Report Generator")
st.caption(
    "Upload your Master File and Result Dump to generate the executive PPTX report "
    "and Excel export packs."
)

# ── Upload gate — stop here until both files are provided ────────────────────

if not master_upload or not result_upload:
    missing = []
    if not master_upload:
        missing.append("**Master File**")
    if not result_upload:
        missing.append("**Result / Dump File**")

    st.warning(
        f"Please upload {' and '.join(missing)} using the sidebar to continue."
    )

    st.markdown("""
    **How to use:**
    1. Upload the **Master File** (MyC Report Excel — Details sheet)
    2. Upload the **Result / Dump File** (assessment proficiency dump)
    3. Select a **Business Group** scope
    4. Generate your PPTX report or download Excel export packs
    """)
    st.stop()

# ── Load data ─────────────────────────────────────────────────────────────────

with st.spinner("Processing files…"):
    try:
        merged_df, resource_df, no_asmt_df = build_data(
            master_upload.getvalue(),
            result_upload.getvalue(),
            selected_bg,
        )
    except Exception as e:
        st.error("Failed to process the uploaded files. Please check the file format.")
        st.exception(e)
        st.stop()

st.success(
    f"✅ Loaded — **{master_upload.name}** + **{result_upload.name}**"
)

# ── Metrics ───────────────────────────────────────────────────────────────────

total_resources    = resource_df[COL_PERSONNEL_NO].nunique()
assessed_resources = resource_df.loc[resource_df["HasAssessment"],  COL_PERSONNEL_NO].nunique()
meeting_target     = resource_df.loc[resource_df["MeetingTarget"],   COL_PERSONNEL_NO].nunique()
below_target       = resource_df.loc[resource_df["BelowTarget"],     COL_PERSONNEL_NO].nunique()
no_assessment      = total_resources - assessed_resources
completion_pct     = assessed_resources / total_resources * 100 if total_resources else 0
compliance_pct     = meeting_target     / total_resources * 100 if total_resources else 0

scope_label = (
    "Tech_Song + Tech_Adobe Platform" if selected_bg == "All" else selected_bg
)

_section(f"{scope_label} — Primary Skills", "📊")

st.markdown(
    '<div class="metric-row">'
    + _card("Total Resources",   f"{total_resources:,}",     accent="blue")
    + _card("Assessed",          f"{assessed_resources:,}",  f"of {total_resources:,}", "blue")
    + _card("Completion",        f"{completion_pct:.1f}%",   "assessed / total",
            "green" if completion_pct >= 80 else "amber")
    + _card("Target Compliance", f"{compliance_pct:.1f}%",   "meeting target / total",
            "green" if compliance_pct >= 80 else "amber")
    + _card("No Assessment",     f"{no_assessment:,}",        "need to complete", "amber")
    + _card("Below Target",      f"{below_target:,}",         "need uplift", "red")
    + "</div>",
    unsafe_allow_html=True,
)

# ── Supporting dataframes ─────────────────────────────────────────────────────

# Career summary
resource_df["career_level_num"] = pd.to_numeric(resource_df["ManagementLevel"], errors="coerce")
career_summary = (
    resource_df
    .groupby("career_level_num", as_index=False)
    .agg(
        TotalResources   =(COL_PERSONNEL_NO, "nunique"),
        AssessedResources=("HasAssessment",  "sum"),
        MeetingTarget    =("MeetingTarget",  "sum"),
        BelowTarget      =("BelowTarget",    "sum"),
    )
    .sort_values("career_level_num", ascending=False)
)
career_summary["No Assessment"] = (
    career_summary["TotalResources"] - career_summary["AssessedResources"]
)
career_summary["Completion %"] = (
    career_summary["AssessedResources"] / career_summary["TotalResources"] * 100
).fillna(0).round(1)
career_summary["Target Compliance %"] = (
    career_summary["MeetingTarget"] / career_summary["TotalResources"] * 100
).fillna(0).round(1)
career_summary = career_summary.rename(columns={
    "career_level_num": "Career Level",
    "TotalResources":   "Total Resources",
    "BelowTarget":      "Below Target",
})

# Skill summary
skill_summary_df = (
    merged_df
    .groupby([COL_SKILL_NAME, COL_SKILL_TYPE], as_index=False)
    .agg(
        Total_Resources   =(COL_PERSONNEL_NO, "nunique"),
        Assessed_Resources=(COL_PERSONNEL_NO,
            lambda x: merged_df.loc[x.index, "has_assessment"].eq(True)
                                .groupby(x).any().sum()),
        Meeting_Target    =(COL_PERSONNEL_NO,
            lambda x: merged_df.loc[x.index, "meets_target"].eq(True)
                                .groupby(x).any().sum()),
        Below_Target      =(COL_PERSONNEL_NO,
            lambda x: merged_df.loc[x.index, "below_target"].eq(True)
                                .groupby(x).any().sum()),
    )
)
skill_summary_df["No_Assessment"] = (
    skill_summary_df["Total_Resources"] - skill_summary_df["Assessed_Resources"]
)
skill_summary_df["Completion %"] = (
    skill_summary_df["Assessed_Resources"] / skill_summary_df["Total_Resources"] * 100
).round(1)
skill_summary_df["Compliance %"] = (
    skill_summary_df["Meeting_Target"] / skill_summary_df["Total_Resources"] * 100
).round(1)
skill_summary_df = skill_summary_df.sort_values("Total_Resources", ascending=False)

# Skill gap for exports
skill_gap_df = skill_summary_df.rename(columns={
    COL_SKILL_NAME:    "Primary Skill",
    "Total_Resources": "Total Resources",
    "No_Assessment":   "No Assessment",
    "Below_Target":    "Below Target Only",
}).copy()
skill_gap_df["Resources To Chase"] = (
    skill_gap_df["No Assessment"] + skill_gap_df["Below Target Only"]
)

# Project summary
project_df = (
    resource_df
    .groupby("Project", as_index=False)
    .agg(
        TotalResources=(COL_PERSONNEL_NO, "nunique"),
        NoAssessment  =("HasAssessment", lambda x: (~x.astype(bool)).sum()),
        BelowTarget   =("BelowTarget",   "sum"),
    )
)
project_df["Resources To Chase"] = project_df["NoAssessment"] + project_df["BelowTarget"]
project_df["Chase %"] = (
    project_df["Resources To Chase"] / project_df["TotalResources"] * 100
).fillna(0).round(1)
project_df = project_df.sort_values("Resources To Chase", ascending=False)

# Resource export
identifier_col = "EID" if show_eid else COL_PERSONNEL_NO
resource_export = resource_df.rename(columns={
    "SkillName":    "Primary Skill",
    "ActionReason": "Action Reason",
}).copy()
resource_export["Career Level"] = pd.to_numeric(
    resource_df["ManagementLevel"], errors="coerce"
)
desired_cols = [identifier_col, "Project", "Primary Skill", "Career Level",
                "TargetProficiency", "ActualProficiency", "Action Reason"]
resource_export = resource_export[[c for c in desired_cols if c in resource_export.columns]]
if "Action Reason" in resource_export.columns:
    _ao = {"No Assessment": 1, "Below Target": 2, "Meeting Target": 3}
    resource_export["_as"] = resource_export["Action Reason"].map(_ao).fillna(99)
    resource_export["_cl"] = pd.to_numeric(resource_export.get("Career Level"), errors="coerce")
    resource_export = resource_export.sort_values(
        ["Project", "_as", "_cl"], na_position="last"
    ).drop(columns=["_as", "_cl"], errors="ignore")

# Summary row
summary_df = pd.DataFrame([{
    "Business Group":      scope_label,
    "Assessment Scope":    "Primary",
    "Total Resources":     total_resources,
    "Assessed Resources":  assessed_resources,
    "Completion %":        round(completion_pct, 1),
    "Target Compliance %": round(compliance_pct, 1),
    "No Assessment":       no_assessment,
    "Below Target":        below_target,
}])

# ── PPTX Generation ───────────────────────────────────────────────────────────

st.markdown("---")
_section("AI Executive Report (PPTX)", "✦")
st.caption(
    "Generates an AI-written executive deck using Azure OpenAI. "
    "Configure `LLM_API_KEY`, `AZURE_OPENAI_ENDPOINT`, and `AZURE_OPENAI_DEPLOYMENT` "
    "in your Streamlit secrets or environment variables."
)

_az_key = _get_secret("LLM_API_KEY") or _get_secret("AZURE_OPENAI_API_KEY")

if not _az_key:
    st.warning(
        "Azure OpenAI credentials not found. "
        "Add `LLM_API_KEY`, `AZURE_OPENAI_ENDPOINT`, and `AZURE_OPENAI_DEPLOYMENT` "
        "to your Streamlit secrets (`.streamlit/secrets.toml`) or environment variables."
    )
else:
    if st.button("✦ Generate AI Executive Report (PPTX)", type="primary"):
        with st.spinner("AI is writing and building the deck…"):
            try:
                # Inject secrets into environment for report_generator
                os.environ["LLM_API_KEY"]              = _get_secret("LLM_API_KEY")
                os.environ["AZURE_OPENAI_ENDPOINT"]    = _get_secret("AZURE_OPENAI_ENDPOINT")
                os.environ["AZURE_OPENAI_DEPLOYMENT"]  = _get_secret("AZURE_OPENAI_DEPLOYMENT")

                from report_generator import build_executive_pptx

                _prof_order = ["P0", "P1", "P2", "P3", "Expert eligible", "Unspecified"]
                _prof_counts = (
                    resource_df["ActualProficiency"].fillna("Unspecified")
                    .value_counts().reindex(_prof_order, fill_value=0)
                )
                _prof_total = _prof_counts.sum() or 1
                proficiency_rows = [
                    {
                        "Proficiency": lbl,
                        "Count":   int(_prof_counts[lbl]),
                        "Percent": round(int(_prof_counts[lbl]) / _prof_total * 100, 1),
                    }
                    for lbl in _prof_order if _prof_counts[lbl] > 0
                ]

                career_rows = career_summary[
                    ["Career Level", "Total Resources", "No Assessment",
                     "Completion %", "Target Compliance %", "Below Target"]
                ].to_dict("records") if not career_summary.empty else []

                project_rows = (
                    project_df[["Project", "Resources To Chase", "Chase %"]]
                    .to_dict("records") if not project_df.empty else []
                )

                if not career_summary.empty:
                    career_summary["_gap"] = (
                        career_summary["Completion %"] - career_summary["Target Compliance %"]
                    )
                    _gap_row  = career_summary.loc[career_summary["_gap"].idxmax()]
                    _priority = career_summary.sort_values("Target Compliance %").iloc[0]
                    career_kpis = {
                        "avg_completion":   float(career_summary["Completion %"].mean()),
                        "avg_compliance":   float(career_summary["Target Compliance %"].mean()),
                        "widest_gap":       int(_gap_row["_gap"]),
                        "priority_segment": f"CL{int(_priority['Career Level'])}",
                    }
                else:
                    career_kpis = {}

                metrics = {
                    "total_resources":    total_resources,
                    "assessed_resources": assessed_resources,
                    "completion_pct":     completion_pct,
                    "compliance_pct":     compliance_pct,
                    "no_assessment":      no_assessment,
                    "below_target":       below_target,
                    "completion_delta": "", "completion_delta_pos": True,
                    "compliance_delta": "", "compliance_delta_pos": True,
                    "no_assessment_delta": "", "no_assessment_delta_pos": True,
                    "below_target_delta": "", "below_target_delta_pos": True,
                }

                pptx_buf = build_executive_pptx(
                    business_group=scope_label,
                    metrics=metrics,
                    career_rows=career_rows,
                    skill_rows=skill_gap_df.to_dict("records"),
                    project_rows=project_rows,
                    proficiency_rows=proficiency_rows,
                    career_kpis=career_kpis,
                )

                fname = (
                    f"mycompetency_executive_report_"
                    f"{scope_label.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.pptx"
                )
                st.success("Report ready!")
                st.download_button(
                    "⬇ Download Executive Report (PPTX)",
                    data=pptx_buf,
                    file_name=fname,
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )

            except Exception as e:
                st.error(f"Report generation failed: {e}")
                st.exception(e)

# ── Excel Exports ──────────────────────────────────────────────────────────────

st.markdown("---")
_section("Excel Exports", "📥")

col1, col2, col3 = st.columns(3)

with col1:
    st.markdown("**Skill Chase Workbook**")
    st.caption("One sheet per skill — No Assessment & Below Target resources.")
    if st.button("Build Skill Chase Workbook"):
        with st.spinner("Building…"):
            try:
                excel = build_skill_chase_workbook(skill_summary_df, merged_df, resource_df)
                st.download_button(
                    "⬇ Download",
                    data=excel,
                    file_name="myCompetency_Skill_Chase_Workbook.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_skill_chase",
                )
            except Exception as e:
                st.error(f"Failed: {e}")

with col2:
    st.markdown("**People Lead Follow-up Pack**")
    st.caption("Executive summary, project actions, resource detail, skill gaps, career health.")
    if st.button("Build People Lead Pack"):
        with st.spinner("Building…"):
            try:
                excel = build_people_lead_excel(
                    summary_df=summary_df,
                    project_df=project_df,
                    resource_df=resource_export,
                    skill_gap_df=skill_gap_df,
                    career_df=career_summary[[
                        "Career Level", "Total Resources", "No Assessment",
                        "Completion %", "Target Compliance %", "Below Target"
                    ]],
                )
                st.download_button(
                    "⬇ Download",
                    data=excel,
                    file_name="mycompetency_people_lead_pack.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_people_lead",
                )
            except Exception as e:
                st.error(f"Failed: {e}")

with col3:
    st.markdown("**No Assessment Action List**")
    st.caption(f"{len(no_asmt_df):,} people — sorted by Details action required.")
    if st.button("Build No Assessment List"):
        with st.spinner("Building…"):
            try:
                excel = build_no_assessment_excel(no_asmt_df)
                st.download_button(
                    "⬇ Download",
                    data=excel,
                    file_name="myCompetency_NoAssessment_List.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_no_asmt",
                )
            except Exception as e:
                st.error(f"Failed: {e}")

# ── No Assessment preview ──────────────────────────────────────────────────────

st.markdown("---")
_section("No Assessment List Preview", "🔍")
st.caption("All No Assessment people — Primary Skills, sorted by Details action required.")
st.dataframe(no_asmt_df, use_container_width=True, hide_index=True)
